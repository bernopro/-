#!/usr/bin/env python3
"""
make_list.py

Usage:
  python make_list.py "path/to/file.docx"

This script extracts lines and table cells from a DOCX and writes audio/list.json
in the project `audio/` folder. If a filename has no extension the script appends
`.mp3`.

Requires: python-docx (install with `python -m pip install python-docx`)
"""
import sys
import json
import os
from docx import Document


def extract(docx_path):
    doc = Document(docx_path)
    texts = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            for line in t.splitlines():
                if line.strip():
                    texts.append(line.strip())
    # also check tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for line in cell.text.splitlines():
                    if line.strip():
                        texts.append(line.strip())
    return texts


def sanitize(name):
    name = name.strip()
    # remove surrounding quotes
    name = name.strip('"').strip("'")
    # if no extension, assume .mp3
    if not os.path.splitext(name)[1]:
        name = name + '.mp3'
    # replace problematic characters for filenames
    name = name.replace('\r','').replace('\n','').strip()
    return name


def main():
    if len(sys.argv) < 2:
        print('Usage: python make_list.py "file.docx"')
        return
    docx_path = sys.argv[1]
    if not os.path.exists(docx_path):
        print('DOCX not found:', docx_path)
        return

    project_root = os.path.dirname(os.path.abspath(__file__))
    audio_dir = os.path.join(project_root, 'audio')
    if not os.path.isdir(audio_dir):
        print('audio/ folder not found in project root:', audio_dir)
        return

    items = extract(docx_path)
    files = [sanitize(i) for i in items if i]
    # dedupe preserving order
    seen = set(); uniq = []
    for f in files:
        if f not in seen:
            seen.add(f); uniq.append(f)

    out_path = os.path.join(audio_dir, 'list.json')
    with open(out_path, 'w', encoding='utf-8') as fh:
        json.dump(uniq, fh, ensure_ascii=False, indent=2)
    print(f'Wrote {out_path} with {len(uniq)} entries.')

    # report missing
    missing = [f for f in uniq if not os.path.exists(os.path.join(audio_dir, f))]
    if missing:
        print('\nMissing files in audio/:')
        for m in missing:
            print('  ', m)
    else:
        print('\nAll files present in audio/')


if __name__ == '__main__':
    main()
